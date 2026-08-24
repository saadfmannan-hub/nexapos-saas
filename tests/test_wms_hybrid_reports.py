"""Focused UI and report coverage for WMS Hybrid compensation."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from django.urls import reverse
from openpyxl import load_workbook

from apps.wms_workforce.models import (
    WmsEmployee,
    WmsEmployeeCategoryAssignment,
)
from tests.test_wms_phase2 import make_employee
from tests.test_wms_phase7 import WmsPhase7Base


class WmsHybridSalaryReportTests(WmsPhase7Base):
    production_date = date(2026, 7, 12)

    def setUp(self):
        super().setUp()
        self.hybrid_employee = make_employee(
            self.business_a,
            self.location_a1,
            "P7-HYBRID-REPORT",
            compensation_type=WmsEmployee.CompensationType.HYBRID,
            fixed_salary=Decimal("200.000"),
            piece_rate=Decimal("0.500"),
        )
        self.hybrid_override = WmsEmployeeCategoryAssignment.objects.create(
            business=self.business_a,
            employee=self.hybrid_employee,
            category=self.category_override,
            per_piece_rate=Decimal("0.700"),
        )
        self.hybrid_default = WmsEmployeeCategoryAssignment.objects.create(
            business=self.business_a,
            employee=self.hybrid_employee,
            category=self.category_default,
            per_piece_rate=None,
        )

    def create_hybrid_production(self):
        return self.create_production(
            employee=self.hybrid_employee,
            production_date=self.production_date,
            daily_total=160,
            override_quantity=100,
            default_quantity=60,
        )

    def calculate_hybrid(self):
        return self.calculate(self.hybrid_employee)

    def test_detail_shows_components_attendance_and_production_breakdowns(self):
        attendance = self.create_attendance(
            employee=self.hybrid_employee,
            attendance_date=self.production_date,
        )
        production = self.create_hybrid_production()
        salary = self.calculate_hybrid()

        self.assertEqual(salary.fixed_salary_component, Decimal("200.000"))
        self.assertEqual(
            salary.production_salary_component,
            Decimal("100.000"),
        )
        self.assertEqual(salary.gross_salary, Decimal("300.000"))

        response = self.client.get(reverse("wms:salary_detail", args=[salary.public_id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Compensation Type")
        self.assertContains(response, "Hybrid (Fixed + Production)")
        self.assertContains(response, "Fixed Salary")
        self.assertContains(response, "Production Earnings")
        self.assertContains(response, "Gross salary")
        self.assertContains(response, "200.000")
        self.assertContains(response, "100.000")
        self.assertContains(response, "300.000")
        self.assertContains(response, "Daily attendance breakdown")
        self.assertContains(response, "No deduction")
        self.assertContains(response, "Override Category")
        self.assertContains(response, "Default Category")
        self.assertContains(response, "Assignment Override")
        self.assertContains(response, "Employee Default")

        attendance_days = response.context["attendance_days"]
        production_days = response.context["production_days"]
        self.assertEqual(len(attendance_days), 31)
        self.assertEqual(len(production_days), 1)

        merged_day = next(day for day in attendance_days if day.salary_date == self.production_date)
        self.assertEqual(merged_day.attendance_id, attendance.pk)
        self.assertEqual(merged_day.production_entry_id, production.pk)
        self.assertEqual(merged_day.eligible_quantity, 160)
        self.assertEqual(merged_day.daily_amount, Decimal("100.000"))
        self.assertEqual(production_days[0].pk, merged_day.pk)

        lines = {line.category_code_snapshot: line for line in merged_day.piece_lines.all()}
        self.assertEqual(lines["OVERRIDE"].quantity, 100)
        self.assertEqual(lines["OVERRIDE"].applied_rate, Decimal("0.700"))
        self.assertEqual(lines["OVERRIDE"].line_amount, Decimal("70.000"))
        self.assertEqual(lines["DEFAULT"].quantity, 60)
        self.assertEqual(lines["DEFAULT"].applied_rate, Decimal("0.500"))
        self.assertEqual(lines["DEFAULT"].line_amount, Decimal("30.000"))

    def test_detail_zero_production_shows_fixed_component_and_empty_state(self):
        salary = self.calculate_hybrid()

        self.assertEqual(salary.fixed_salary_component, Decimal("200.000"))
        self.assertEqual(
            salary.production_salary_component,
            Decimal("0.000"),
        )
        self.assertEqual(salary.gross_salary, Decimal("200.000"))

        response = self.client.get(reverse("wms:salary_detail", args=[salary.public_id]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["attendance_days"]), 31)
        self.assertEqual(response.context["production_days"], [])
        self.assertContains(response, "Hybrid (Fixed + Production)")
        self.assertContains(response, "Fixed Salary")
        self.assertContains(response, "Production Earnings")
        self.assertContains(
            response,
            "No eligible production exists for this month.",
        )
        self.assertNotContains(response, "No production category lines.")

    def test_salary_list_shows_hybrid_components_and_gross_total(self):
        self.create_hybrid_production()
        salary = self.calculate_hybrid()

        response = self.client.get(
            reverse("wms:salary_list"),
            {"month": "2026-07"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["record_count"], 1)
        listed_salary = response.context["page"].object_list[0]
        self.assertEqual(listed_salary.pk, salary.pk)
        self.assertContains(response, "Fixed component")
        self.assertContains(response, "Production earnings")
        self.assertContains(response, "Gross salary")
        self.assertContains(response, "200.000")
        self.assertContains(response, "100.000")
        self.assertContains(response, "300.000")

    def test_salary_report_filter_returns_exact_hybrid_components(self):
        self.create_hybrid_production()
        hybrid_salary = self.calculate_hybrid()
        self.calculate(self.fixed_employee)
        self.calculate(self.piece_employee)

        response = self.client.get(
            reverse("wms:report_salary"),
            {
                "report_month": "2026-07",
                "salary_type": WmsEmployee.CompensationType.HYBRID,
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["form"].is_valid())
        report = response.context["report"]
        self.assertEqual(len(report["rows"]), 1)
        row = report["rows"][0]
        self.assertEqual(row["employee_code"], self.hybrid_employee.employee_code)
        self.assertEqual(row["salary_type"], "Hybrid (Fixed + Production)")
        self.assertEqual(row["base_salary"], Decimal("200.000"))
        self.assertEqual(row["eligible_pieces"], 160)
        self.assertEqual(row["piece_earnings"], Decimal("100.000"))
        self.assertEqual(row["final_salary"], Decimal("300.000"))
        self.assertEqual(row["public_id"], hybrid_salary.public_id)
        self.assertEqual(report["totals"]["base_salary"], Decimal("200.000"))
        self.assertEqual(report["totals"]["eligible_pieces"], 160)
        self.assertEqual(
            report["totals"]["piece_earnings"],
            Decimal("100.000"),
        )
        self.assertEqual(report["totals"]["final_salary"], Decimal("300.000"))
        self.assertContains(response, "Fixed Component")
        self.assertContains(response, "Production Earnings")
        self.assertContains(response, "Gross Salary")
        self.assertContains(response, "Print / Save PDF")

    def test_salary_xlsx_has_hybrid_component_headings_and_values(self):
        self.create_hybrid_production()
        salary = self.calculate_hybrid()

        response = self.client.get(
            reverse("wms:report_salary_export"),
            {
                "report_month": "2026-07",
                "salary_type": WmsEmployee.CompensationType.HYBRID,
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        header_index = next(
            index
            for index, row in enumerate(rows)
            if row[:3] == ("Employee", "Employee Code", "Salary Type")
        )
        expected_headings = (
            "Employee",
            "Employee Code",
            "Salary Type",
            "Fixed Component",
            "Eligible Production Pieces",
            "Production Earnings",
            "Gross Salary",
            "Calculation Status",
            "Calculated / Last Updated",
        )
        self.assertEqual(rows[header_index][:9], expected_headings)

        values = rows[header_index + 1]
        self.assertEqual(values[0], self.hybrid_employee.full_name)
        self.assertEqual(values[1], self.hybrid_employee.employee_code)
        self.assertEqual(values[2], "Hybrid (Fixed + Production)")
        self.assertEqual(Decimal(str(values[3])), Decimal("200.000"))
        self.assertEqual(values[4], 160)
        self.assertEqual(Decimal(str(values[5])), Decimal("100.000"))
        self.assertEqual(Decimal(str(values[6])), Decimal("300.000"))
        self.assertEqual(values[7], salary.get_status_display())

        totals = rows[header_index + 2]
        self.assertEqual(totals[0], "Totals")
        self.assertEqual(Decimal(str(totals[3])), Decimal("200.000"))
        self.assertEqual(totals[4], 160)
        self.assertEqual(Decimal(str(totals[5])), Decimal("100.000"))
        self.assertEqual(Decimal(str(totals[6])), Decimal("300.000"))

        self.assertEqual(sheet.cell(header_index + 2, 4).number_format, "#,##0.000")
        self.assertEqual(sheet.cell(header_index + 2, 6).number_format, "#,##0.000")
        self.assertEqual(sheet.cell(header_index + 2, 7).number_format, "#,##0.000")
