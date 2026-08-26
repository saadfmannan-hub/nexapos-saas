"""Regression coverage for expense drawer linking and monthly reporting."""
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import Membership, Role, User
from apps.branches.models import Branch
from apps.expenses import services as expense_services
from apps.expenses.models import Expense, ExpenseCategory
from apps.registers import services as register_services
from apps.registers.models import CashRegister, Shift
from apps.sales.models import PaymentMethod
from apps.subscriptions.exceptions import ModuleAccessDenied

from .base import TenantTestCase


D = Decimal


class ExpenseRegisterUpgradeTests(TenantTestCase):
    def setUp(self):
        self.client.force_login(self.owner_a)
        self.category = ExpenseCategory.objects.for_business(self.business_a).first()
        self.bank = PaymentMethod.objects.for_business(self.business_a).get(kind="bank")
        self.online = PaymentMethod.objects.create(
            business=self.business_a,
            name="Online Test",
            kind=PaymentMethod.Kind.ONLINE,
        )

    def make_expense(self, *, number, amount="10.000", **overrides):
        values = {
            "business": self.business_a,
            "expense_number": number,
            "expense_date": timezone.localdate(),
            "branch": self.branch_a,
            "category": self.category,
            "amount": D(amount),
            "status": Expense.Status.APPROVED,
            "created_by": self.owner_a,
        }
        values.update(overrides)
        return Expense.objects.create(**values)

    def expense_payload(self, **overrides):
        values = {
            "expense_date": timezone.localdate().isoformat(),
            "branch": str(self.branch_a.pk),
            "category": str(self.category.pk),
            "payee": "Drawer supplier",
            "supplier": "",
            "amount": "10.000",
            "tax_amount": "0.000",
            "payment_method": str(self.cash_a.pk),
            "reference": "UPGRADE",
            "description": "Expense upgrade test",
        }
        values.update(overrides)
        return values

    def open_shift(self, opening="41.000"):
        return register_services.open_shift(
            business=self.business_a,
            register=self.register_a,
            cashier=self.owner_a,
            opening_cash=D(opening),
            membership=self.membership_a(),
        )

    def make_staff(self, *permissions):
        role = Role.objects.create(
            business=self.business_a,
            name=f"Expense register role {User.objects.count()}",
            permissions=list(permissions),
        )
        user = User.objects.create_user(
            email=f"expense-register-{User.objects.count()}@example.com",
            password="StrongPass123!",
            full_name="Expense Register Staff",
        )
        membership = Membership.objects.create(
            business=self.business_a,
            user=user,
            role=role,
        )
        membership.branches.set([self.branch_a])
        return user, membership

    def make_staff_shift(self, user, opening="20.000"):
        return Shift.objects.create(
            business=self.business_a,
            register=self.register_a,
            branch=self.branch_a,
            cashier=user,
            opened_at=timezone.now(),
            opening_cash=D(opening),
        )

    def close_staff_shift(self, shift, actual="15.000"):
        return register_services.close_shift(
            shift=shift,
            actual_cash=D(actual),
            user=self.owner_a,
            membership=self.membership_a(),
        )

    def test_production_cash_expense_example_reconciles_to_zero(self):
        shift = self.open_shift()
        self.make_expense(
            number="EXP-PRODUCTION",
            amount="34.400",
            payment_method=self.cash_a,
            shift=shift,
        )

        totals = register_services.shift_totals(shift)
        self.assertEqual(totals["cash_expenses"], D("34.400"))
        self.assertEqual(totals["expected_cash"], D("6.600"))

        register_services.close_shift(
            shift=shift,
            actual_cash=D("6.600"),
            user=self.owner_a,
            notes="Drawer balanced after documented expense.",
        )
        shift.refresh_from_db()
        self.assertEqual(shift.expected_cash, D("6.600"))
        self.assertEqual(shift.actual_cash, D("6.600"))
        self.assertEqual(shift.difference, D("0.000"))

    def test_only_exact_linked_cash_reduces_the_drawer(self):
        shift = self.open_shift("100.000")
        cases = (
            ("UNLINKED-CASH", self.cash_a, None),
            ("CARD", self.card_a, shift),
            ("BANK", self.bank, shift),
            ("ONLINE", self.online, shift),
        )
        for suffix, method, linked_shift in cases:
            self.make_expense(
                number=f"EXP-{suffix}",
                amount="5.000",
                payment_method=method,
                shift=linked_shift,
            )
        totals = register_services.shift_totals(shift)
        self.assertEqual(totals["cash_expenses"], D("0.000"))
        self.assertEqual(totals["expected_cash"], D("100.000"))

    def test_register_branch_and_tenant_expenses_are_isolated(self):
        shift_a = self.open_shift("50.000")
        register_two = CashRegister.objects.create(
            business=self.business_a,
            branch=self.branch_a,
            name="Register Two",
            code="REG-TWO",
        )
        shift_two = Shift.objects.create(
            business=self.business_a,
            register=register_two,
            branch=self.branch_a,
            cashier=self.cashier_a,
            opened_at=timezone.now(),
            opening_cash=D("20.000"),
        )
        self.make_expense(
            number="EXP-REGISTER-TWO",
            payment_method=self.cash_a,
            shift=shift_two,
        )

        other_branch = Branch.objects.create(
            business=self.business_a,
            name="Other Sales Branch",
            code="OTHER-SALES",
        )
        other_register = CashRegister.objects.create(
            business=self.business_a,
            branch=other_branch,
            name="Other Branch Register",
            code="OTHER-REG",
        )
        other_shift = Shift.objects.create(
            business=self.business_a,
            register=other_register,
            branch=other_branch,
            cashier=self.cashier_a,
            opened_at=timezone.now(),
            opening_cash=D("10.000"),
        )
        self.make_expense(
            number="EXP-OTHER-BRANCH",
            branch=other_branch,
            payment_method=self.cash_a,
            shift=other_shift,
        )

        category_b = ExpenseCategory.objects.for_business(self.business_b).first()
        cash_b = PaymentMethod.objects.for_business(self.business_b).get(kind="cash")
        Expense.objects.create(
            business=self.business_b,
            expense_number="EXP-OTHER-TENANT",
            expense_date=timezone.localdate(),
            branch=self.branch_b,
            category=category_b,
            amount=D("10.000"),
            payment_method=cash_b,
            shift=shift_a,
        )

        self.assertEqual(
            register_services.shift_totals(shift_a)["cash_expenses"],
            D("0.000"),
        )

    def test_create_requires_payment_medium_and_links_only_matching_cash_drawer(self):
        shift = self.open_shift()
        missing = self.client.post(
            reverse("expenses:create"),
            self.expense_payload(payment_method=""),
        )
        self.assertEqual(missing.status_code, 200)
        self.assertContains(missing, "This field is required")

        linked = self.client.post(
            reverse("expenses:create"),
            self.expense_payload(
                amount="34.400",
                reference="DRAWER-LINKED",
                paid_from_drawer="on",
            ),
        )
        self.assertRedirects(linked, reverse("expenses:list"))
        expense = Expense.objects.get(reference="DRAWER-LINKED")
        self.assertEqual(expense.shift_id, shift.pk)
        self.assertEqual(expense.created_by_id, self.owner_a.pk)

        non_cash = self.client.post(
            reverse("expenses:create"),
            self.expense_payload(
                payment_method=str(self.bank.pk),
                reference="BANK-NO-DRAWER",
                paid_from_drawer="on",
            ),
        )
        self.assertRedirects(non_cash, reverse("expenses:list"))
        self.assertIsNone(Expense.objects.get(reference="BANK-NO-DRAWER").shift_id)

        other_branch = Branch.objects.create(
            business=self.business_a,
            name="Mismatch Branch",
            code="MISMATCH",
        )
        mismatch = self.client.post(
            reverse("expenses:create"),
            self.expense_payload(
                branch=str(other_branch.pk),
                reference="BRANCH-NO-DRAWER",
                paid_from_drawer="on",
            ),
        )
        self.assertRedirects(mismatch, reverse("expenses:list"))
        self.assertIsNone(Expense.objects.get(reference="BRANCH-NO-DRAWER").shift_id)

    def test_service_rejects_cross_tenant_or_cross_branch_shift(self):
        shift = self.open_shift()
        other_branch = Branch.objects.create(
            business=self.business_a,
            name="Service Mismatch",
            code="SERVICE-MISMATCH",
        )
        expense = Expense(
            business=self.business_a,
            expense_number="EXP-INVALID-SERVICE",
            expense_date=timezone.localdate(),
            branch=other_branch,
            category=self.category,
            amount=D("1.000"),
            payment_method=self.cash_a,
            created_by=self.owner_a,
        )
        with self.assertRaises(ValidationError):
            expense_services.save_manual_expense(
                expense=expense,
                business=self.business_a,
                user=self.owner_a,
                membership=self.membership_a(),
                requested_shift=shift,
                drawer_requested=True,
            )
        self.assertFalse(
            Expense.objects.filter(expense_number="EXP-INVALID-SERVICE").exists()
        )

    def test_non_cash_change_clears_a_stale_drawer_link(self):
        shift = self.open_shift()
        expense = self.make_expense(
            number="EXP-STALE-DRAWER",
            payment_method=self.cash_a,
            shift=shift,
        )
        expense.payment_method = self.bank

        expense_services.save_manual_expense(
            expense=expense,
            business=self.business_a,
            user=self.owner_a,
            membership=self.membership_a(),
            requested_shift=shift,
            drawer_requested=True,
        )

        expense.refresh_from_db()
        self.assertIsNone(expense.shift_id)

    def test_open_shift_financial_edit_needs_only_expense_manage(self):
        user, _membership = self.make_staff("expenses.manage")
        shift = self.make_staff_shift(user)
        expense = self.make_expense(
            number="EXP-OPEN-EDIT",
            amount="5.000",
            status=Expense.Status.SUBMITTED,
            payment_method=self.cash_a,
            shift=shift,
            created_by=user,
        )
        self.client.force_login(user)

        response = self.client.post(
            reverse("expenses:edit", args=[expense.public_id]),
            self.expense_payload(amount="6.000", paid_from_drawer="on"),
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("expenses:list"))
        expense.refresh_from_db()
        self.assertEqual(expense.amount, D("6.000"))
        self.assertEqual(expense.shift_id, shift.pk)

    def test_closed_shift_amount_edit_requires_shift_approval(self):
        user, _membership = self.make_staff("expenses.manage")
        shift = self.make_staff_shift(user)
        expense = self.make_expense(
            number="EXP-CLOSED-DENIED",
            amount="5.000",
            status=Expense.Status.SUBMITTED,
            payment_method=self.cash_a,
            shift=shift,
            created_by=user,
        )
        shift = self.close_staff_shift(shift)
        self.client.force_login(user)

        response = self.client.post(
            reverse("expenses:edit", args=[expense.public_id]),
            self.expense_payload(amount="6.000"),
        )

        self.assertEqual(response.status_code, 403)
        expense.refresh_from_db()
        shift.refresh_from_db()
        self.assertEqual(expense.amount, D("5.000"))
        self.assertEqual(expense.shift_id, shift.pk)
        self.assertEqual(shift.expected_cash, D("15.000"))
        self.assertEqual(shift.difference, D("0.000"))

    def test_closed_shift_amount_edit_with_shift_approval_refreshes_snapshot(self):
        user, _membership = self.make_staff(
            "expenses.manage",
            "shifts.approve",
        )
        shift = self.make_staff_shift(user)
        expense = self.make_expense(
            number="EXP-CLOSED-ALLOWED",
            amount="5.000",
            status=Expense.Status.SUBMITTED,
            payment_method=self.cash_a,
            shift=shift,
            created_by=user,
        )
        shift = self.close_staff_shift(shift)
        self.client.force_login(user)

        response = self.client.post(
            reverse("expenses:edit", args=[expense.public_id]),
            self.expense_payload(amount="6.000"),
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("expenses:list"))
        expense.refresh_from_db()
        shift.refresh_from_db()
        self.assertEqual(expense.amount, D("6.000"))
        self.assertEqual(expense.shift_id, shift.pk)
        self.assertEqual(shift.expected_cash, D("14.000"))
        self.assertEqual(shift.difference, D("1.000"))

    def test_closed_shift_cash_to_non_cash_edit_requires_shift_approval(self):
        user, _membership = self.make_staff("expenses.manage")
        shift = self.make_staff_shift(user)
        expense = self.make_expense(
            number="EXP-CLOSED-NON-CASH",
            amount="5.000",
            status=Expense.Status.SUBMITTED,
            payment_method=self.cash_a,
            shift=shift,
            created_by=user,
        )
        shift = self.close_staff_shift(shift)
        self.client.force_login(user)

        response = self.client.post(
            reverse("expenses:edit", args=[expense.public_id]),
            self.expense_payload(payment_method=str(self.bank.pk)),
        )

        self.assertEqual(response.status_code, 403)
        expense.refresh_from_db()
        shift.refresh_from_db()
        self.assertEqual(expense.payment_method_id, self.cash_a.pk)
        self.assertEqual(expense.shift_id, shift.pk)
        self.assertEqual(shift.expected_cash, D("15.000"))

    def test_closed_shift_link_remove_or_change_requires_shift_approval(self):
        user, membership = self.make_staff("expenses.manage")
        shift = self.make_staff_shift(user)
        expense = self.make_expense(
            number="EXP-CLOSED-LINK",
            amount="5.000",
            status=Expense.Status.SUBMITTED,
            payment_method=self.cash_a,
            shift=shift,
            created_by=user,
        )
        shift = self.close_staff_shift(shift)

        with self.assertRaises(ModuleAccessDenied):
            expense_services.save_manual_expense(
                expense=Expense.objects.get(pk=expense.pk),
                business=self.business_a,
                user=user,
                membership=membership,
                requested_shift=None,
                drawer_requested=False,
            )

        replacement = self.make_staff_shift(user)
        with self.assertRaises(ModuleAccessDenied):
            expense_services.save_manual_expense(
                expense=Expense.objects.get(pk=expense.pk),
                business=self.business_a,
                user=user,
                membership=membership,
                requested_shift=replacement,
                drawer_requested=True,
            )

        expense.refresh_from_db()
        shift.refresh_from_db()
        self.assertEqual(expense.shift_id, shift.pk)
        self.assertEqual(shift.expected_cash, D("15.000"))

    def test_closed_shift_status_qualification_change_requires_shift_approval(self):
        user, membership = self.make_staff("expenses.approve")
        shift = self.make_staff_shift(user)
        expense = self.make_expense(
            number="EXP-CLOSED-STATUS",
            amount="5.000",
            status=Expense.Status.SUBMITTED,
            payment_method=self.cash_a,
            shift=shift,
            created_by=user,
        )
        shift = self.close_staff_shift(shift)

        with self.assertRaises(ModuleAccessDenied):
            expense_services.set_expense_status(
                expense=expense,
                status=Expense.Status.REJECTED,
                user=user,
                membership=membership,
            )

        expense.refresh_from_db()
        shift.refresh_from_db()
        self.assertEqual(expense.status, Expense.Status.SUBMITTED)
        self.assertEqual(shift.expected_cash, D("15.000"))

    def test_closed_shift_metadata_only_edit_needs_no_shift_approval(self):
        user, _membership = self.make_staff("expenses.manage")
        shift = self.make_staff_shift(user)
        expense = self.make_expense(
            number="EXP-CLOSED-METADATA",
            amount="5.000",
            status=Expense.Status.SUBMITTED,
            payment_method=self.cash_a,
            shift=shift,
            created_by=user,
        )
        shift = self.close_staff_shift(shift)
        self.client.force_login(user)

        response = self.client.post(
            reverse("expenses:edit", args=[expense.public_id]),
            self.expense_payload(
                amount="5.000",
                payee="Updated payee only",
                description="Updated metadata only",
            ),
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("expenses:list"))
        expense.refresh_from_db()
        shift.refresh_from_db()
        self.assertEqual(expense.payee, "Updated payee only")
        self.assertEqual(expense.description, "Updated metadata only")
        self.assertEqual(expense.shift_id, shift.pk)
        self.assertEqual(shift.expected_cash, D("15.000"))
        self.assertEqual(shift.difference, D("0.000"))

    def test_historical_correction_still_requires_both_approval_permissions(self):
        user, membership = self.make_staff(
            "expenses.manage",
            "expenses.approve",
        )
        expense = self.make_expense(
            number="EXP-HISTORICAL-PERMISSION",
            amount="5.000",
            payment_method=self.cash_a,
            created_by=user,
        )
        expense.payment_method = self.bank

        with self.assertRaises(ModuleAccessDenied):
            expense_services.save_manual_expense(
                expense=expense,
                business=self.business_a,
                user=user,
                membership=membership,
                requested_shift=None,
                drawer_requested=False,
                historical_correction=True,
            )

        expense.refresh_from_db()
        self.assertEqual(expense.payment_method_id, self.cash_a.pk)

    def test_ambiguous_open_drawers_are_not_auto_linked(self):
        self.open_shift()
        second_register = CashRegister.objects.create(
            business=self.business_a,
            branch=self.branch_a,
            name="Second Current Drawer",
            code="SECOND-CURRENT",
        )
        Shift.objects.create(
            business=self.business_a,
            register=second_register,
            branch=self.branch_a,
            cashier=self.owner_a,
            opened_at=timezone.now(),
            opening_cash=D("5.000"),
        )

        response = self.client.post(
            reverse("expenses:create"),
            self.expense_payload(
                reference="AMBIGUOUS-DRAWER",
                paid_from_drawer="on",
            ),
        )

        self.assertRedirects(response, reverse("expenses:list"))
        self.assertIsNone(
            Expense.objects.get(reference="AMBIGUOUS-DRAWER").shift_id
        )

    def test_historical_relink_is_audited_without_duplication_and_refreshes_close(self):
        shift = self.open_shift()
        register_services.close_shift(
            shift=shift,
            actual_cash=D("6.600"),
            user=self.owner_a,
            notes="Preserve this note",
        )
        expense = self.make_expense(
            number="EXP-HISTORICAL",
            amount="34.400",
            payment_method=self.cash_a,
        )
        count_before = Expense.objects.count()

        response = self.client.post(
            reverse("expenses:edit", args=[expense.public_id]),
            {
                "payment_method": str(self.cash_a.pk),
                "historical_shift": str(shift.pk),
            },
        )
        self.assertRedirects(response, reverse("expenses:list"))
        expense.refresh_from_db()
        shift.refresh_from_db()
        self.assertEqual(Expense.objects.count(), count_before)
        self.assertEqual(expense.shift_id, shift.pk)
        self.assertEqual(shift.expected_cash, D("6.600"))
        self.assertEqual(shift.difference, D("0.000"))
        self.assertEqual(shift.actual_cash, D("6.600"))
        self.assertEqual(shift.closing_notes, "Preserve this note")
        self.assertTrue(
            self.business_a.audit_logs.filter(
                action="expense.drawer_corrected"
            ).exists()
        )

    def test_reopen_and_reclose_counts_linked_expense_once(self):
        shift = self.open_shift("20.000")
        self.make_expense(
            number="EXP-RECLOSE",
            amount="5.000",
            payment_method=self.cash_a,
            shift=shift,
        )
        register_services.close_shift(
            shift=shift, actual_cash=D("15.000"), user=self.owner_a
        )
        register_services.reopen_shift(shift=shift, user=self.owner_a)
        register_services.close_shift(
            shift=shift, actual_cash=D("15.000"), user=self.owner_a
        )
        shift.refresh_from_db()
        self.assertEqual(shift.expected_cash, D("15.000"))
        self.assertEqual(shift.difference, D("0.000"))
        self.assertEqual(Expense.objects.filter(expense_number="EXP-RECLOSE").count(), 1)

    def test_rejected_and_cancelled_expenses_remain_excluded(self):
        shift = self.open_shift("30.000")
        for status, number in (
            (Expense.Status.APPROVED, "EXP-EFFECTIVE"),
            (Expense.Status.REJECTED, "EXP-REJECTED"),
            (Expense.Status.CANCELLED, "EXP-CANCELLED"),
        ):
            self.make_expense(
                number=number,
                amount="5.000",
                status=status,
                payment_method=self.cash_a,
                shift=shift,
            )
        self.assertEqual(
            register_services.shift_totals(shift)["cash_expenses"],
            D("5.000"),
        )

    def test_expense_amount_remains_total_drawer_outflow_not_amount_plus_tax(self):
        shift = self.open_shift("20.000")
        self.make_expense(
            number="EXP-TAX-SEMANTICS",
            amount="10.000",
            tax_amount=D("0.500"),
            payment_method=self.cash_a,
            shift=shift,
        )
        self.assertEqual(
            register_services.shift_totals(shift)["expected_cash"],
            D("10.000"),
        )

    def test_closing_notes_render_on_detail_and_printable_page(self):
        shift = self.open_shift()
        register_services.close_shift(
            shift=shift,
            actual_cash=D("41.000"),
            user=self.owner_a,
            notes="Closing note visible in print",
        )
        response = self.client.get(
            reverse("registers:shift_detail", args=[shift.public_id])
        )
        self.assertContains(response, "Closing note visible in print")
        self.assertContains(response, "z-report-closing-notes")
        self.assertNotContains(
            response,
            'z-report-closing-notes no-print',
        )

    def test_list_paid_via_filter_legacy_label_and_three_decimal_display(self):
        self.business_a.currency_precision = 3
        self.business_a.save(update_fields=["currency_precision"])
        shift = self.open_shift()
        self.make_expense(
            number="EXP-LIST-DRAWER",
            amount="34.400",
            payee="EXP-LIST-DRAWER",
            payment_method=self.cash_a,
            shift=shift,
        )
        self.make_expense(
            number="EXP-LIST-UNSPECIFIED",
            amount="1.300",
            payee="EXP-LIST-UNSPECIFIED",
            payment_method=None,
        )
        response = self.client.get(reverse("expenses:list"))
        self.assertContains(response, "Paid Via")
        self.assertContains(response, "Cash – Register")
        self.assertContains(response, "Unspecified")
        self.assertContains(response, "34.400")

        filtered = self.client.get(
            reverse("expenses:list"), {"method": "cash_register"}
        )
        self.assertContains(filtered, "EXP-LIST-DRAWER")
        self.assertNotContains(filtered, "EXP-LIST-UNSPECIFIED")


class MonthlyExpenseReportUpgradeTests(TenantTestCase):
    def setUp(self):
        self.client.force_login(self.owner_a)
        self.business_a.currency_precision = 3
        self.business_a.save(update_fields=["currency_precision"])
        self.category = ExpenseCategory.objects.for_business(self.business_a).first()
        self.bank = PaymentMethod.objects.for_business(self.business_a).get(kind="bank")

    def make_expense(self, number, day, amount, method=None, **overrides):
        values = {
            "business": self.business_a,
            "expense_number": number,
            "expense_date": date(2026, 8, day),
            "branch": self.branch_a,
            "category": self.category,
            "payee": number,
            "amount": D(amount),
            "payment_method": method,
            "status": Expense.Status.APPROVED,
            "created_by": self.owner_a,
        }
        values.update(overrides)
        return Expense.objects.create(**values)

    def test_monthly_summary_payment_category_and_every_day_are_correct(self):
        self.make_expense("EXP-AUG-CASH", 2, "10.000", self.cash_a)
        self.make_expense("EXP-AUG-CARD", 2, "20.400", self.card_a)
        self.make_expense("EXP-AUG-BANK", 25, "57.300", self.bank)
        self.make_expense("EXP-JULY", 31, "999.000", self.cash_a).expense_date = date(2026, 7, 31)
        Expense.objects.filter(expense_number="EXP-JULY").update(expense_date=date(2026, 7, 31))

        response = self.client.get(
            reverse("reports:view", args=["expense_analysis"]),
            {"month": "2026-08"},
        )
        self.assertEqual(response.status_code, 200)
        data = response.context["data"]
        self.assertEqual(data["total_expenses"], D("87.700"))
        self.assertEqual(data["transaction_count"], 3)
        self.assertEqual(data["average_expense"], D("29.233"))
        self.assertEqual(len(data["daily_breakdown"]), 31)
        august_second = data["daily_breakdown"][1]
        self.assertEqual(august_second["count"], 2)
        self.assertEqual(august_second["total"], D("30.400"))
        self.assertEqual(data["daily_breakdown"][24]["total"], D("57.300"))
        self.assertEqual(data["rows"][0][1], 3)
        payment = {row["key"]: row for row in data["payment_breakdown"]}
        self.assertEqual(payment["cash_non_register"]["total"], D("10.000"))
        self.assertEqual(payment["card"]["total"], D("20.400"))
        self.assertEqual(payment["bank"]["total"], D("57.300"))
        self.assertContains(response, "87.700")
        self.assertContains(response, "30.400")
        self.assertContains(response, "Day-wise Monthly Breakdown")
        self.assertContains(response, "Daily Transaction Detail")

    def test_monthly_branch_payment_status_and_tenant_filters_are_isolated(self):
        other_branch = Branch.objects.create(
            business=self.business_a,
            name="Monthly Other",
            code="MONTH-OTHER",
        )
        self.make_expense("EXP-FILTER-CASH", 3, "5.000", self.cash_a)
        self.make_expense(
            "EXP-FILTER-BANK", 4, "7.000", self.bank, branch=other_branch
        )
        self.make_expense(
            "EXP-FILTER-CANCEL", 5, "9.000", self.cash_a,
            status=Expense.Status.CANCELLED,
        )
        category_b = ExpenseCategory.objects.for_business(self.business_b).first()
        Expense.objects.create(
            business=self.business_b,
            expense_number="EXP-TENANT-B",
            expense_date=date(2026, 8, 3),
            branch=self.branch_b,
            category=category_b,
            amount=D("100.000"),
            payment_method=PaymentMethod.objects.for_business(self.business_b).get(
                kind="cash"
            ),
        )

        response = self.client.get(
            reverse("reports:view", args=["expense_analysis"]),
            {
                "month": "2026-08",
                "branch": str(self.branch_a.pk),
                "method": "cash_non_register",
            },
        )
        self.assertEqual(response.context["data"]["total_expenses"], D("5.000"))
        self.assertEqual(response.context["data"]["transaction_count"], 1)

        cancelled = self.client.get(
            reverse("reports:view", args=["expense_analysis"]),
            {"month": "2026-08", "status": Expense.Status.CANCELLED},
        )
        self.assertEqual(cancelled.context["data"]["total_expenses"], D("9.000"))

    def test_monthly_xlsx_has_three_sheets_and_pdf_receives_all_sections(self):
        self.make_expense("EXP-EXPORT", 8, "6.600", self.cash_a)
        url = reverse("reports:view", args=["expense_analysis"])
        xlsx = self.client.get(url, {"month": "2026-08", "export": "xlsx"})
        self.assertEqual(xlsx.status_code, 200)
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(xlsx.content), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["Monthly Summary", "Daily Breakdown", "Expense Transactions"],
        )
        workbook.close()

        with patch(
            "apps.reports.exports.render_pdf", return_value=b"%PDF-test"
        ) as render_pdf:
            pdf = self.client.get(url, {"month": "2026-08", "export": "pdf"})
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        template_name, context = render_pdf.call_args.args
        self.assertEqual(template_name, "reports/expense_analysis_pdf.html")
        self.assertEqual(context["data"]["total_expenses"], D("6.600"))
        self.assertEqual(len(context["data"]["daily_breakdown"]), 31)
